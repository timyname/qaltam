from __future__ import annotations

from dataclasses import dataclass
import csv
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
import io
from pathlib import Path
import re
from typing import Any
import unicodedata

from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.localization import normalize_language, text
from backend.app.models.enums import TransactionType
from backend.app.models.profile import ImportedStatement


DATE_PATTERN = re.compile(r"(?P<date>\d{2}[./-]\d{2}[./-]\d{2,4}|\d{4}-\d{2}-\d{2})")
AMOUNT_PATTERN = re.compile(r"(?P<amount>[+-]?\d[\d\s']*(?:[.,]\d{1,2})?)")

DATE_COLUMNS = (
    "date",
    "transaction date",
    "operation date",
    "posted date",
    "value date",
    "\u0434\u0430\u0442\u0430",
    "\u0434\u0430\u0442\u0430 \u043e\u043f\u0435\u0440\u0430\u0446\u0438\u0438",
    "\u0434\u0430\u0442\u0430 \u043f\u0440\u043e\u0432\u043e\u0434\u043a\u0438",
    "\u0434\u0430\u0442\u0430 \u0432\u0430\u043b\u044e\u0442\u0438\u0440\u043e\u0432\u0430\u043d\u0438\u044f",
)
AMOUNT_COLUMNS = (
    "amount",
    "sum",
    "total",
    "amount kzt",
    "\u0441\u0443\u043c\u043c\u0430",
    "\u0441\u0443\u043c\u043c\u0430 kzt",
    "\u0441\u0443\u043c\u043c\u0430 \u043e\u043f\u0435\u0440\u0430\u0446\u0438\u0438",
)
DEBIT_COLUMNS = (
    "debit",
    "withdrawal",
    "expense",
    "\u0440\u0430\u0441\u0445\u043e\u0434",
    "\u0441\u043f\u0438\u0441\u0430\u043d\u0438\u0435",
    "\u0434\u0435\u0431\u0435\u0442",
)
CREDIT_COLUMNS = (
    "credit",
    "deposit",
    "income",
    "\u043f\u0440\u0438\u0445\u043e\u0434",
    "\u0437\u0430\u0447\u0438\u0441\u043b\u0435\u043d\u0438\u0435",
    "\u043a\u0440\u0435\u0434\u0438\u0442",
)
COUNTERPARTY_COLUMNS = (
    "counterparty",
    "merchant",
    "payee",
    "receiver",
    "sender",
    "client",
    "\u043a\u043e\u043d\u0442\u0440\u0430\u0433\u0435\u043d\u0442",
    "\u043f\u043e\u043b\u0443\u0447\u0430\u0442\u0435\u043b\u044c",
    "\u043e\u0442\u043f\u0440\u0430\u0432\u0438\u0442\u0435\u043b\u044c",
    "\u043f\u043b\u0430\u0442\u0435\u043b\u044c\u0449\u0438\u043a",
)
DESCRIPTION_COLUMNS = (
    "description",
    "details",
    "purpose",
    "comment",
    "\u043d\u0430\u0437\u043d\u0430\u0447\u0435\u043d\u0438\u0435",
    "\u043e\u043f\u0438\u0441\u0430\u043d\u0438\u0435",
    "\u043d\u0430\u0437\u043d\u0430\u0447\u0435\u043d\u0438\u0435 \u043f\u043b\u0430\u0442\u0435\u0436\u0430",
)

CSV_DELIMITERS = (",", ";", "\t", "|")
TEXT_FILE_ENCODINGS = ("utf-8-sig", "utf-16", "cp1251", "koi8-r")


@dataclass(slots=True)
class StatementEntry:
    statement_date: date
    amount: Decimal
    transaction_type: TransactionType
    counterparty: str
    description: str
    source_line: str


@dataclass(slots=True)
class ParsedStatement:
    imported_statement: ImportedStatement
    entries: list[StatementEntry]
    raw_content: str


class StatementParser:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def parse_file(
        self,
        file_path: Path,
        *,
        telegram_user_id: int,
        original_filename: str | None = None,
        language: str = "ru",
    ) -> ParsedStatement:
        language = normalize_language(language, default="ru")
        raw_content = self._read_raw_content(file_path)
        imported_statement = ImportedStatement(
            source_name=self._detect_source_name(original_filename or file_path.name, raw_content),
            file_path=str(file_path),
            original_filename=original_filename or file_path.name,
            telegram_user_id=telegram_user_id,
            raw_content=raw_content,
            parse_status="pending",
        )
        self.session.add(imported_statement)
        await self.session.flush()

        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            entries = self._parse_tabular_file(file_path, reader="csv", language=language)
        elif suffix in {".xlsx", ".xls"}:
            entries = self._parse_tabular_file(file_path, reader="excel", language=language)
        elif suffix == ".pdf":
            entries = self._parse_text_lines(raw_content.splitlines(), language=language)
        else:
            raise ValueError(f"Unsupported statement format: {suffix}")

        imported_statement.parse_status = "parsed" if entries else "empty"
        imported_statement.note = self._build_parse_note(
            len(entries),
            imported_statement.source_name,
            language=language,
        )
        await self.session.flush()
        return ParsedStatement(imported_statement=imported_statement, entries=entries, raw_content=raw_content)

    async def parse_text(
        self,
        raw_text: str,
        *,
        telegram_user_id: int,
        source_name: str = "telegram_text_statement",
        language: str = "ru",
    ) -> ParsedStatement:
        language = normalize_language(language, default="ru")
        synthetic_path = (
            f"telegram://statement/{telegram_user_id}/"
            f"{datetime.now(UTC).strftime('%Y%m%d%H%M%S%f')}.txt"
        )
        imported_statement = ImportedStatement(
            source_name=self._detect_source_name(source_name, raw_text),
            file_path=synthetic_path,
            original_filename=f"{source_name}.txt",
            telegram_user_id=telegram_user_id,
            raw_content=raw_text,
            parse_status="pending",
        )
        self.session.add(imported_statement)
        await self.session.flush()

        entries = self._parse_text_lines(raw_text.splitlines(), language=language)
        imported_statement.parse_status = "parsed" if entries else "empty"
        imported_statement.note = self._build_parse_note(len(entries), self._text_statement_source_name(language), language=language)
        await self.session.flush()
        return ParsedStatement(imported_statement=imported_statement, entries=entries, raw_content=raw_text)

    @staticmethod
    def looks_like_statement_text(text: str) -> bool:
        normalized = text.strip()
        if not normalized:
            return False
        if normalized.count("\n") >= 1:
            return len(DATE_PATTERN.findall(normalized)) >= 2 and len(AMOUNT_PATTERN.findall(normalized)) >= 2
        return StatementParser._looks_like_statement_row(normalized)

    @staticmethod
    def _looks_like_statement_row(text: str) -> bool:
        line = " ".join(text.strip().split())
        if not line:
            return False
        date_match = DATE_PATTERN.search(line)
        if not date_match:
            return False
        if not StatementParser._parse_date(date_match.group("date")):
            return False
        remainder = (line[: date_match.start()] + " " + line[date_match.end() :]).strip()
        amount_matches = list(AMOUNT_PATTERN.finditer(remainder))
        if not amount_matches:
            return False
        return StatementParser._parse_decimal(amount_matches[-1].group("amount")) is not None

    def _read_raw_content(self, file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            try:
                import pdfplumber
            except Exception as exc:  # pragma: no cover - optional dependency
                raise RuntimeError("pdfplumber is required for PDF statement parsing.") from exc
            texts: list[str] = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    if page_text:
                        texts.append(page_text)
            return "\n".join(texts).strip()
        if suffix == ".csv":
            return self._read_text_file(file_path)
        if suffix in {".xlsx", ".xls"}:
            try:
                from openpyxl import load_workbook
            except Exception as exc:  # pragma: no cover - optional dependency
                raise RuntimeError("openpyxl is required for Excel statement parsing.") from exc
            workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                return ""
            headers = [str(cell or "") for cell in rows[0]]
            content_lines = [",".join(headers)]
            for row in rows[1:]:
                content_lines.append(",".join(str(cell or "") for cell in row))
            return "\n".join(content_lines)
        return self._read_text_file(file_path)

    def _parse_tabular_file(self, file_path: Path, *, reader: str, language: str) -> list[StatementEntry]:
        if reader == "csv":
            return self._parse_csv_rows(file_path, language=language)
        return self._parse_excel_rows(file_path, language=language)

    def _parse_text_lines(self, lines: list[str], *, language: str) -> list[StatementEntry]:
        entries: list[StatementEntry] = []
        for raw_line in lines:
            line = " ".join(raw_line.strip().split())
            if not line:
                continue
            date_match = DATE_PATTERN.search(line)
            if not date_match:
                continue
            statement_date = self._parse_date(date_match.group("date"))
            if not statement_date:
                continue
            remainder = (line[: date_match.start()] + " " + line[date_match.end() :]).strip()
            amount_matches = list(AMOUNT_PATTERN.finditer(remainder))
            if not amount_matches:
                continue
            amount_match = amount_matches[-1]
            amount_value = self._parse_decimal(amount_match.group("amount"))
            if amount_value is None:
                continue
            description = (remainder[: amount_match.start()] + " " + remainder[amount_match.end() :]).strip(" |-")
            transaction_type = self._infer_transaction_type(amount_match.group("amount"), description)
            entries.append(
                StatementEntry(
                    statement_date=statement_date,
                    amount=abs(amount_value),
                    transaction_type=transaction_type,
                    counterparty=self._derive_counterparty(description),
                    description=description or self._statement_item_label(language),
                    source_line=line,
                )
            )
        return entries

    def _extract_amount(
        self,
        amount_value: Any,
        debit_value: Any,
        credit_value: Any,
    ) -> tuple[Decimal, TransactionType] | None:
        direct_amount = self._parse_decimal(amount_value)
        if direct_amount is not None:
            transaction_type = TransactionType.INCOME if direct_amount >= 0 else TransactionType.EXPENSE
            return abs(direct_amount), transaction_type

        debit_amount = self._parse_decimal(debit_value)
        if debit_amount is not None and debit_amount != 0:
            return abs(debit_amount), TransactionType.EXPENSE

        credit_amount = self._parse_decimal(credit_value)
        if credit_amount is not None and credit_amount != 0:
            return abs(credit_amount), TransactionType.INCOME
        return None

    @staticmethod
    def _pick_column(columns: dict[Any, str], aliases: tuple[str, ...]) -> Any | None:
        for original, normalized in columns.items():
            if normalized in aliases:
                return original
        for original, normalized in columns.items():
            if any(alias in normalized for alias in aliases):
                return original
        return None

    @staticmethod
    def _normalize_key(value: Any) -> str:
        text = unicodedata.normalize("NFKC", str(value)).strip().casefold()
        return " ".join("".join(char if char.isalnum() else " " for char in text).split())

    @staticmethod
    def _parse_date(value: Any) -> date | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        if not text:
            return None
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
        except ValueError:
            pass
        for fmt in (
            "%d.%m.%Y",
            "%d.%m.%y",
            "%d/%m/%Y",
            "%d/%m/%y",
            "%d-%m-%Y",
            "%d-%m-%y",
            "%Y-%m-%d",
            "%d.%m.%Y %H:%M",
            "%d.%m.%Y %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d %H:%M:%S",
        ):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _parse_decimal(value: Any) -> Decimal | None:
        if value is None:
            return None
        text = str(value).strip()
        if not text:
            return None
        is_negative = text.startswith("(") and text.endswith(")")
        cleaned = (
            text.replace("\u00a0", "")
            .replace("\u2009", "")
            .replace("\u202f", "")
            .replace("\u2212", "-")
            .replace("\u2014", "-")
            .replace("\u20b8", "")
            .replace("KZT", "")
            .replace("kzt", "")
            .replace("'", "")
            .replace(" ", "")
        )
        if "," in cleaned and "." in cleaned:
            if cleaned.rfind(",") > cleaned.rfind("."):
                cleaned = cleaned.replace(".", "").replace(",", ".")
            else:
                cleaned = cleaned.replace(",", "")
        else:
            cleaned = cleaned.replace(",", ".")
        cleaned = re.sub(r"[^0-9.+-]", "", cleaned)
        if is_negative and cleaned and not cleaned.startswith("-"):
            cleaned = f"-{cleaned.lstrip('+')}"
        if cleaned in {"", "+", "-", "."}:
            return None
        try:
            return Decimal(cleaned)
        except (InvalidOperation, ValueError):
            return None

    @staticmethod
    def _string_value(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        return "" if text.lower() == "nan" else text

    @staticmethod
    def _infer_transaction_type(raw_amount: str, description: str) -> TransactionType:
        lowered = description.lower()
        if raw_amount.strip().startswith("+"):
            return TransactionType.INCOME
        if raw_amount.strip().startswith("-"):
            return TransactionType.EXPENSE
        if any(
            token in lowered
            for token in (
                "\u043f\u043e\u0441\u0442\u0443\u043f",
                "\u0437\u0430\u0447\u0438\u0441",
                "\u043f\u0440\u0438\u0445\u043e\u0434",
                "income",
                "receipt",
                "client",
                "avans",
            )
        ):
            return TransactionType.INCOME
        return TransactionType.EXPENSE

    @staticmethod
    def _derive_counterparty(description: str) -> str:
        cleaned = description.replace("|", " ").replace(";", " ").strip()
        tokens = [token for token in cleaned.split() if token]
        if not tokens:
            return "Unknown"
        return " ".join(tokens[:3])[:80]

    @staticmethod
    def _detect_source_name(filename: str, raw_content: str) -> str:
        lowered = f"{filename} {raw_content[:2000]}".lower()
        if "kaspi" in lowered:
            return "Kaspi"
        if "halyk" in lowered or "\u043d\u0430\u0440\u043e\u0434\u043d" in lowered:
            return "Halyk"
        if "statement" in lowered:
            return "Bank Statement"
        return "Imported Statement"

    def _parse_csv_rows(self, file_path: Path, *, language: str) -> list[StatementEntry]:
        csv_text = self._read_text_file(file_path)
        sample = "\n".join(csv_text.splitlines()[:8])
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters="".join(CSV_DELIMITERS))
            reader_kwargs: dict[str, Any] = {"dialect": dialect, "skipinitialspace": True}
        except csv.Error:
            reader_kwargs = {"delimiter": self._fallback_csv_delimiter(csv_text), "skipinitialspace": True}

        with io.StringIO(csv_text, newline="") as handle:
            reader = csv.DictReader(handle, **reader_kwargs)
            return self._parse_row_dicts(list(reader), language=language)

    def _parse_excel_rows(self, file_path: Path, *, language: str) -> list[StatementEntry]:
        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("openpyxl is required for Excel statement parsing.") from exc

        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell or "") for cell in rows[0]]
        row_dicts: list[dict[str, Any]] = []
        for row in rows[1:]:
            row_dicts.append({headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))})
        return self._parse_row_dicts(row_dicts, language=language)

    def _parse_row_dicts(self, rows: list[dict[str, Any]], *, language: str) -> list[StatementEntry]:
        if not rows:
            return []
        normalized_columns = {column: self._normalize_key(column) for column in rows[0].keys()}
        date_column = self._pick_column(normalized_columns, DATE_COLUMNS)
        amount_column = self._pick_column(normalized_columns, AMOUNT_COLUMNS)
        debit_column = self._pick_column(normalized_columns, DEBIT_COLUMNS)
        credit_column = self._pick_column(normalized_columns, CREDIT_COLUMNS)
        counterparty_column = self._pick_column(normalized_columns, COUNTERPARTY_COLUMNS)
        description_column = self._pick_column(normalized_columns, DESCRIPTION_COLUMNS)

        entries: list[StatementEntry] = []
        for row in rows:
            raw_date = row.get(date_column) if date_column else None
            statement_date = self._parse_date(raw_date)
            amount_info = self._extract_amount(
                row.get(amount_column) if amount_column else None,
                row.get(debit_column) if debit_column else None,
                row.get(credit_column) if credit_column else None,
            )
            if not statement_date or not amount_info:
                continue

            amount, transaction_type = amount_info
            counterparty = self._string_value(row.get(counterparty_column)) if counterparty_column else ""
            description = self._string_value(row.get(description_column)) if description_column else ""
            if not description and counterparty:
                description = counterparty
            if not counterparty and description:
                counterparty = self._derive_counterparty(description)
            if not description and not counterparty:
                description = self._statement_item_label(language)

            source_line = " | ".join(self._string_value(value) for value in row.values() if self._string_value(value))
            entries.append(
                StatementEntry(
                    statement_date=statement_date,
                    amount=amount,
                    transaction_type=transaction_type,
                    counterparty=counterparty or self._unknown_counterparty_label(language),
                    description=description or counterparty or self._statement_item_label(language),
                    source_line=source_line,
                )
            )
        return entries

    @staticmethod
    def _fallback_csv_delimiter(csv_text: str) -> str:
        first_line = next((line for line in csv_text.splitlines() if line.strip()), "")
        counts = {delimiter: first_line.count(delimiter) for delimiter in CSV_DELIMITERS}
        delimiter, count = max(counts.items(), key=lambda item: item[1])
        return delimiter if count > 0 else ","

    @staticmethod
    def _read_text_file(file_path: Path) -> str:
        raw_bytes = file_path.read_bytes()
        for encoding in TEXT_FILE_ENCODINGS:
            try:
                return raw_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw_bytes.decode("utf-8", errors="ignore")

    @staticmethod
    def _statement_item_label(language: str) -> str:
        return "Элемент выписки" if normalize_language(language, default="ru") == "ru" else "Statement item"

    @staticmethod
    def _unknown_counterparty_label(language: str) -> str:
        return "Неизвестно" if normalize_language(language, default="ru") == "ru" else "Unknown"

    @staticmethod
    def _text_statement_source_name(language: str) -> str:
        return "текстовой выписки" if normalize_language(language, default="ru") == "ru" else "text statement"

    @staticmethod
    def _build_parse_note(parsed_count: int, source_name: str, *, language: str) -> str:
        if normalize_language(language, default="ru") == "ru":
            return f"Разобрано {parsed_count} строк из {source_name}"
        return text("parsed_items", language, count=parsed_count)
